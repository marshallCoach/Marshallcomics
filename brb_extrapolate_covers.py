#!/usr/bin/env python3
"""brb_extrapolate_covers.py — turn confirmed Fandom covers into leads for the
rest of each series.

Once a book has a Fandom-sourced cover, its (Title, Volume) has a known-good
Fandom page pattern "{Title} Vol {vol} {issue}". This finds every OTHER issue of
that same Title+Volume that is still missing a cover and writes a constructed
Fandom page URL for it — a data-fixes "fandom" record. Feed the output to
brb_apply_fandom_pages.py, which fetches each page and writes the cover.

Run it after every cover apply: each pass confirms more series, so the pool
compounds. Read-only on the xlsx + covers.json.

    python3 brb_extrapolate_covers.py
    python3 brb_apply_fandom_pages.py --flags fandom-extrapolated-<date>.json --apply
    python3 brb.py --commit "apply extrapolated sibling covers" --yes
"""
import glob, json, os, re, datetime

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "attached_assets")
COVERS = os.path.join(ROOT, "covers.json")


def latest_xlsx():
    m = [f for f in glob.glob(os.path.join(ASSETS, "comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        raise SystemExit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=os.path.getmtime)


def ni(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def nv(v):
    try:
        return str(int(float(str(v).strip())))
    except (ValueError, TypeError):
        return "1"


def main():
    import openpyxl
    cov = json.load(open(COVERS))

    def U(e):
        return (e.get("url") if isinstance(e, dict) else e) if e else None

    def covered(t, iss, vol):
        return any(U(cov.get(k)) for k in
                   (f"{t}|||{iss}|||{vol}", f"{t}|||{iss}", f"{t}|||#{iss}"))

    # Confirmed (title.lower, vol) -> wiki, from any Fandom-sourced cover.
    conf = {}
    for k, e in cov.items():
        if not (isinstance(e, dict) and U(e)):
            continue
        src = str(e.get("source", ""))
        if not src.startswith("fandom"):
            continue
        p = k.split("|||")
        if len(p) >= 3 and p[2].strip():
            conf[(p[0].strip().lower(), p[2].strip())] = "dc" if "dc" in src else "marvel"

    src = latest_xlsx()
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("✅ Clean Inventory"))
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0]); C = {n: H.index(n) for n in H if n}

    def g(r, n):
        i = C.get(n); return r[i] if i is not None else None

    out, seen = [], set()
    now = datetime.datetime.now().isoformat()
    for r in rows[1:]:
        t = str(g(r, "Title") or "").strip()
        iss = ni(g(r, "Issue #")); vol = nv(g(r, "Volume"))
        box = str(g(r, "Box #") or "").strip()
        if not t or not iss:
            continue
        key = (t.lower(), vol)
        if key not in conf or covered(t, iss, vol):
            continue
        dd = (t.lower(), iss, vol)
        if dd in seen:
            continue
        seen.add(dd)
        wiki = conf[key]
        url = f"https://{wiki}.fandom.com/wiki/{t.replace(' ', '_')}_Vol_{vol}_{iss}"
        out.append({"id": f"{t}|||{iss}|||{box}", "title": t, "issue": iss, "box": box,
                    "problem": "no-cover", "kind": "fandom", "value": url, "at": now})

    today = datetime.date.today().isoformat()
    path = os.path.join(ROOT, f"fandom-extrapolated-{today}.json")
    json.dump(out, open(path, "w"), indent=2, ensure_ascii=False)
    print(f"Source: {os.path.basename(src)}")
    print(f"Confirmed Fandom series+volumes: {len(conf)}")
    print(f"Extrapolated missing-cover leads: {len(out)} -> {os.path.basename(path)}")
    if out:
        print(f"Apply: python3 brb_apply_fandom_pages.py --flags {os.path.basename(path)} --apply")


if __name__ == "__main__":
    main()
