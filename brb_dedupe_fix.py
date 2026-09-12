#!/usr/bin/env python3
"""brb_dedupe_fix.py — apply the 4 reviewed duplicate decisions.

Reviewed with volume data (per the "always check volume" rule):
  KILL (true duplicates):
    - Entry 12290  Aliens vs. Avengers #1 (2024) — reversed-title dup of 12291
    - Entry 9363   New Fantastic Four #1 (2022) — dup of 9362 (Peter David mini)
  KEEP BUT FIX (not duplicates):
    - Entry 10775  Batgirl and the Birds of Prey #1 (2016) — the REBIRTH one-shot,
                   distinct from the ongoing #1 (10776). Tag its notes so it stops
                   flagging as a dup.
    - Entry 13234  Doom #1 (2024) — same signed book as 84 but Volume said 2 vs 3;
                   align it to 3 (the confirmed Sanford Greene SS row).

Matches rows by the '#' (Entry) column — exact, so nothing else is touched.
Dry-run by default; --apply writes a NEW dated xlsx (never edits in place).

    python3 brb_dedupe_fix.py           # show what would change
    python3 brb_dedupe_fix.py --apply
Then: python3 brb.py --commit "resolve reviewed duplicates" --yes
"""
import argparse, datetime, glob, os, sys

ROOT = os.path.dirname(os.path.abspath(__file__))

DELETE_ENTRIES = {"12290", "9363"}
NOTE_TAG = {"10775": "Batgirl and the Birds of Prey: Rebirth #1 (one-shot) — distinct from the ongoing #1."}
VOL_FIX = {"13234": "3"}


def newest_xlsx():
    cands = [f for f in glob.glob(os.path.join(ROOT, "attached_assets/comics_inventory_*.xlsx"))
             if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not cands:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(cands, key=os.path.getmtime)


def norm_entry(v):
    s = str(v or "").strip()
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    import openpyxl
    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    def ci(n):
        return H.index(n) + 1 if n in H else None
    cEntry, cTitle, cIssue, cVol, cNotes = (ci("#"), ci("Title"), ci("Issue #"),
                                            ci("Volume"), ci("Seller Notes / Variants / Caveats"))
    if not cEntry:
        sys.exit("no '#' (Entry) column found")

    to_delete, notes_done, vol_done = [], [], []
    for r in range(2, ws.max_row + 1):
        e = norm_entry(ws.cell(r, cEntry).value)
        title = str(ws.cell(r, cTitle).value or "").strip()
        if e in DELETE_ENTRIES:
            to_delete.append((r, e, title))
        if e in NOTE_TAG and cNotes:
            notes_done.append((r, e, title))
        if e in VOL_FIX and cVol:
            vol_done.append((r, e, title, str(ws.cell(r, cVol).value or "")))

    print(f"SOURCE: {os.path.basename(src)}\n")
    print("KILL (delete rows):")
    for r, e, t in to_delete:
        print(f"  row {r}  Entry {e}  {t}")
    print("\nTAG NOTES (keep, mark distinct):")
    for r, e, t in notes_done:
        print(f"  row {r}  Entry {e}  {t}  ->  {NOTE_TAG[e]!r}")
    print("\nFIX VOLUME:")
    for r, e, t, cur in vol_done:
        print(f"  row {r}  Entry {e}  {t}  Vol {cur!r} -> {VOL_FIX[e]!r}")

    missing = (DELETE_ENTRIES | set(NOTE_TAG) | set(VOL_FIX)) - {x[1] for x in to_delete} \
        - {x[1] for x in notes_done} - {x[1] for x in vol_done}
    if missing:
        print(f"\n⚠ entries NOT found (already resolved?): {sorted(missing)}")

    if not args.apply:
        print("\nDRY RUN — nothing written. Re-run with --apply to write a new xlsx.")
        return

    for r, e, t, _cur in vol_done:
        ws.cell(r, cVol, VOL_FIX[e])
    for r, e, t in notes_done:
        cur = str(ws.cell(r, cNotes).value or "").strip()
        ws.cell(r, cNotes, (cur + " " + NOTE_TAG[e]).strip() if NOTE_TAG[e] not in cur else cur)
    for r, e, t in sorted(to_delete, reverse=True):   # delete bottom-up
        ws.delete_rows(r, 1)

    out = os.path.join(ROOT, f"attached_assets/comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
    wb.save(out)
    print(f"\nWROTE: {os.path.basename(out)}  ({len(to_delete)} rows deleted, "
          f"{len(notes_done)} notes tagged, {len(vol_done)} volumes fixed)")
    print("Next: python3 brb.py --commit \"resolve reviewed duplicates\" --yes")


if __name__ == "__main__":
    main()
