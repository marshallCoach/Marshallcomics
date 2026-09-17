#!/usr/bin/env python3
"""brb_apply_fandom_pages.py — make a dropped Fandom link true source of truth.

For every 'fandom' resolution in a Data Fix export, this OPENS the linked Fandom
page (not just the URL) and pulls the real publication date + cover, then fills
them for that issue AND its owned run-mates (same title, same volume from the
URL), year-gated so a wrong-era page is never used. Free (Fandom API, no usage $).

Fills:
  - Publication Date column (blank rows only, unless --overwrite)
  - covers.json volume-aware key  Title|||Issue|||Vol
  - Writer(s)/Artist(s)/Cover Artist from the infobox, via the same safe
    fill/union/conflict merge as brb_credits_gcd (never downgrades a name).
    Best free credit source for the not-in-GCD books GCD couldn't reach.

Runs on the Mac. Writes a NEW dated xlsx (never edits in place) + covers.json.
Dry-run by default.

    python3 brb_apply_fandom_pages.py --flags ~/Downloads/data-fixes-merged.json
    python3 brb_apply_fandom_pages.py --flags ~/Downloads/data-fixes-merged.json --apply
Then: python3 brb.py --commit "apply Fandom page data" --yes
"""
import argparse, datetime, glob, json, os, re, sys, time, urllib.parse, urllib.request

from brb_credits_gcd import decide, nk, PLACEHOLDERS  # same safe fill/union/conflict merge

ROOT = os.path.dirname(os.path.abspath(__file__))
COVERS = os.path.join(ROOT, "covers.json")
PUBLIC = os.path.join(ROOT, "artifacts/comics-inventory/public/covers.json")
UA = "BlackReadBrown-Comics/1.0"
DELAY = 0.4
MONTHS = {m.lower(): i for i, m in enumerate(
    ["", "January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"])}


def ni(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def nv(v):
    try:
        return str(int(float(str(v).strip())))
    except (ValueError, TypeError):
        return "1"


def yr(v):
    n = [int(t) for t in re.findall(r"\d{4}", str(v or "")) if 1900 < int(t) < 2100]
    return n[0] if n else None


# Marvel/DC Fandom infobox credit fields. The Marvel Database template uses
# numbered fields (Writer1_1, Penciler1_1, CoverArtist1); some pages use plain
# Writers=/Pencilers=/CoverArtists=. Match both.
CRED_PATS = {
    "writer": r"\|\s*Writers?\d*_?\d*\s*=\s*([^\n|]+)",
    "artist": r"\|\s*(?:Pencilers?|Artists?)\d*_?\d*\s*=\s*([^\n|]+)",
    # Marvel Database credits the COVER via Image1_Artist1/2 (the artists of the
    # primary cover image), not a CoverArtist field. Match both, but only
    # Image1_* (Image2_* is the variant/textless cover).
    "cover":  r"\|\s*(?:Cover\s?Artists?\d*|Image1_Artists?\d*)\s*=\s*([^\n|]+)",
}


def clean_credit(raw):
    """One infobox value -> list of creator names. Unwraps [[Link|Text]] (keeps
    the link target), drops {{templates}} and wiki markup, splits on ; / ,."""
    raw = re.sub(r"\{\{[^}]*\}\}", " ", raw or "")
    raw = re.sub(r"\[\[([^\]|]+)(?:\|[^\]]+)?\]\]", r"\1", raw)  # [[A|B]] -> A
    raw = raw.replace("[[", " ").replace("]]", " ").replace("'''", " ")
    out = []
    for p in re.split(r"\s*(?:;|/|,| and )\s*", raw):
        p = p.strip()
        if p and nk(p) not in PLACEHOLDERS:
            out.append(p)
    return out


def infobox_names(wt, pat):
    """All creator names for a role (dedup, order-preserving)."""
    out, seen = [], set()
    for m in re.findall(pat, wt):
        for nm in clean_credit(m):
            if nk(nm) not in seen:
                seen.add(nk(nm)); out.append(nm)
    return out


def parse_url(u):
    """(api_base, canonical_title, vol) from a .../wiki/Title_Vol_N_ISS URL."""
    m = re.match(r"https?://([^/]+)/wiki/(.+)", u or "")
    if not m:
        return None
    base = f"https://{m.group(1)}/api.php"
    path = urllib.parse.unquote(m.group(2))
    vm = re.search(r"^(.*?)_Vol_(\d+)(?:_|$)", path)  # trailing issue optional
    if not vm:
        return None
    return base, vm.group(1).replace("_", " ").strip(), vm.group(2)


def fetch_page(base, title, vol, issue):
    """Return (image_filename, date_str 'YYYY-MM-00', page_year, credits) or None.
    credits = {'writer':[...], 'artist':[...], 'cover':[...]} from the infobox."""
    u = base + "?" + urllib.parse.urlencode({
        "action": "parse", "page": f"{title} Vol {vol} {issue}",
        "prop": "wikitext", "format": "json", "formatversion": 2, "redirects": 1})
    try:
        d = json.load(urllib.request.urlopen(urllib.request.Request(u, headers={"User-Agent": UA}), timeout=25))
        time.sleep(DELAY)
    except Exception:
        time.sleep(DELAY); return None
    wt = d.get("parse", {}).get("wikitext", "")
    if not wt:
        return None
    im = re.search(r"\|\s*Image1?\s*=\s*([^\n|]+\.(?:jpg|png|jpeg))", wt, re.I)
    dm = re.search(r"\|\s*(?:CoverDate|ReleaseDate)\s*=\s*([^\n|]+)", wt)
    date_str, py = None, None
    if dm:
        raw = dm.group(1)
        py = yr(raw)
        mon = next((MONTHS[w.lower()] for w in re.findall(r"[A-Za-z]+", raw) if w.lower() in MONTHS), 0)
        if py:
            date_str = f"{py}-{mon:02d}-00"
    if not py:
        py = yr(wt[:400])
    credits = {role: infobox_names(wt, pat) for role, pat in CRED_PATS.items()}
    return (im.group(1).strip() if im else None, date_str, py, credits)


def img_url(base, fn):
    try:
        u = base + "?" + urllib.parse.urlencode({
            "action": "query", "titles": "File:" + fn, "prop": "imageinfo",
            "iiprop": "url", "format": "json", "formatversion": 2})
        d = json.load(urllib.request.urlopen(urllib.request.Request(u, headers={"User-Agent": UA}), timeout=25))
        time.sleep(DELAY)
        for p in d.get("query", {}).get("pages", []):
            if "imageinfo" in p:
                return p["imageinfo"][0]["url"].split("/revision/")[0]
    except Exception:
        pass
    return None


def newest_xlsx():
    m = [f for f in glob.glob(os.path.join(ROOT, "attached_assets/comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=os.path.getmtime)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--flags", help="path to a data-fixes export")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--overwrite", action="store_true")
    args = ap.parse_args()

    path = args.flags or max(glob.glob(os.path.expanduser("~/Downloads/data-fixes*.json")) +
                             glob.glob("data-fixes*.json"), key=os.path.getmtime, default=None)
    if not path or not os.path.exists(path):
        sys.exit("no data-fixes export found — pass --flags <path>")
    all_recs = json.load(open(path))
    recs = [r for r in all_recs if r.get("kind") == "fandom" and "_Vol_" in (r.get("value") or "")]
    img_recs = [r for r in all_recs if r.get("kind") == "image" and (r.get("value") or "").startswith("http")]
    print(f"FANDOM links to open: {len(recs)}   direct image links: {len(img_recs)}")

    import openpyxl
    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    def ci(n):
        return H.index(n) + 1 if n in H else None
    cT, cI, cY, cV, cPD = ci("Title"), ci("Issue #"), ci("Year"), ci("Volume"), ci("Publication Date")
    if cPD is None:
        cPD = ws.max_column + 1; ws.cell(1, cPD, "Publication Date")
    CRED_COLS = {"writer": ci("Writer(s)"), "artist": ci("Artist(s)"), "cover": ci("Cover Artist")}

    # index rows by lowercase title
    from collections import defaultdict
    by_title = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        t = str(ws.cell(r, cT).value or "").strip()
        if t:
            by_title[t.lower()].append(r)

    cov = json.load(open(COVERS))
    dates_set = covers_set = 0
    cred_counts = {"writer": 0, "artist": 0, "cover": 0}
    cred_changes = []
    for rec in recs:
        pu = parse_url(rec.get("value"))
        if not pu:
            continue
        base, canon, vol = pu
        sheet_title = rec.get("title") or ""
        for r in by_title.get(sheet_title.lower(), []):
            issue = ni(ws.cell(r, cI).value)
            ry = yr(ws.cell(r, cY).value)
            res = fetch_page(base, canon, vol, issue)
            if not res:
                continue
            fn, date_str, py, credits = res
            if ry and py and abs(py - ry) > 2:   # year-gate: wrong era, skip all
                continue
            if date_str and (args.overwrite or not str(ws.cell(r, cPD).value or "").strip()):
                if args.apply:
                    ws.cell(r, cPD, date_str)
                dates_set += 1
            if fn:
                url = img_url(base, fn)
                key = f"{sheet_title}|||{issue}|||{nv(ws.cell(r, cV).value)}"
                if url and (args.overwrite or key not in cov or not (cov.get(key) or {})):
                    if args.apply:
                        cov[key] = {"url": url, "large": url, "date": str(py or ""), "source": "fandom-page"}
                    covers_set += 1
            # credits — same safe fill/union/conflict merge as the GCD run
            for role, col in CRED_COLS.items():
                if not col:
                    continue
                action, newv = decide(ws.cell(r, col).value, credits.get(role, []))
                if not action:
                    continue
                old = str(ws.cell(r, col).value or "")
                if args.apply:
                    ws.cell(r, col, newv)
                cred_counts[role] += 1
                cred_changes.append([sheet_title, issue, role, action, old, newv])
            print(f"  {sheet_title} #{issue}: date={date_str or '—'} cover={'yes' if fn else 'no'} "
                  f"W/A/C={len(credits['writer'])}/{len(credits['artist'])}/{len(credits['cover'])}")

    # Direct image links (kind="image" from the Cover modal's "Save image") —
    # set the volume-aware covers.json key straight to the pasted URL, no fetch.
    imgs_set = 0
    for rec in img_recs:
        sheet_title = rec.get("title") or ""
        url = (rec.get("value") or "").strip()
        for r in by_title.get(sheet_title.lower(), []):
            if str(rec.get("issue") or "") and ni(ws.cell(r, cI).value) != ni(rec.get("issue")):
                continue
            key = f"{sheet_title}|||{ni(ws.cell(r, cI).value)}|||{nv(ws.cell(r, cV).value)}"
            if args.overwrite or key not in cov or not (cov.get(key) or {}):
                if args.apply:
                    cov[key] = {"url": url, "large": url, "source": "manual-link"}
                imgs_set += 1
                print(f"  image: {sheet_title} #{ni(ws.cell(r, cI).value)} -> {url[:60]}")

    if cred_changes:
        import csv as _csv
        cp = os.path.join(ROOT, "credits_fandom_changes.csv")
        with open(cp, "w", newline="") as f:
            w = _csv.writer(f); w.writerow(["Title", "Issue", "Role", "Action", "Old", "New"]); w.writerows(cred_changes)
        print(f"\nCredit change log: {os.path.basename(cp)} ({len(cred_changes)} rows)")

    print(f"\nDates to set: {dates_set}   Covers to set: {covers_set}   Direct image links: {imgs_set}")
    print(f"Credits — writer {cred_counts['writer']}  artist {cred_counts['artist']}  cover {cred_counts['cover']}")
    if not args.apply:
        print("DRY RUN — nothing written. Re-run with --apply.")
        return
    out = os.path.join(ROOT, f"attached_assets/comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
    wb.save(out)
    json.dump(cov, open(COVERS, "w"))
    if os.path.exists(PUBLIC):
        json.dump(cov, open(PUBLIC, "w"))
    print(f"WROTE: {os.path.basename(out)}  (+ covers.json)")
    print("Next: python3 brb.py --commit \"apply Fandom page data\" --yes")


if __name__ == "__main__":
    main()
