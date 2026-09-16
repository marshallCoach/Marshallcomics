#!/usr/bin/env python3
"""brb_gcd_notfound.py — list every owned comic that GCD has NO series for.

Re-checks the whole sheet against the local GCD catalog (same title+publisher
matching as brb_gcd_volume_check) and emits the books GCD can't place at all —
these are the ones needing a Fandom link or a manual call. Output feeds the
Data Fix page's "Not in GCD" category:

    artifacts/comics-inventory/public/gcd_notfound.json
    ["Title|||Issue|||Box", ...]

Read-only on the xlsx. Run on the Mac (GCD is local). Commit the JSON so the
site picks it up:  python3 brb.py --commit "GCD not-found list" --yes
"""
import glob, json, os, re, sqlite3, sys
from collections import defaultdict
from brb_gcd_volume_check import pub_match, tight

ROOT = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(ROOT, "gcd_local.sqlite")
OUT = os.path.join(ROOT, "artifacts/comics-inventory/public/gcd_notfound.json")


def latest_xlsx():
    m = [f for f in glob.glob(os.path.join(ROOT, "attached_assets/comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=os.path.getmtime)


def norm_issue(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def main():
    if not os.path.exists(DB):
        sys.exit(f"GCD database not found: {DB}")
    conn = sqlite3.connect(DB)
    by_key = defaultdict(list)
    for name, yb, pub in conn.execute(
        "SELECT s.matched_title, s.year_began, p.name "
        "FROM gcd_series s LEFT JOIN gcd_publisher p ON p.id = s.publisher_id"):
        by_key[tight(name)].append({"yb": yb, "pub": pub})
    try:
        aliases = dict(conn.execute("SELECT alias, matched_title FROM gcd_title_alias"))
    except sqlite3.OperationalError:
        aliases = {}

    def has_series(title, pub):
        pool = list(by_key.get(tight(title), []))
        al = aliases.get(title)
        if al:
            pool += by_key.get(tight(al), [])
        return any(pub_match(pub, s["pub"]) for s in pool) if pub else bool(pool)

    import openpyxl
    xlsx = latest_xlsx()
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("✅ Clean Inventory"))
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0])
    ti, ii, pi, bi = (H.index(c) for c in ("Title", "Issue #", "Publisher", "Box #"))

    notfound, total = [], 0
    for r in rows[1:]:
        title = str(r[ti] or "").strip()
        if not title:
            continue
        total += 1
        if not has_series(title, str(r[pi] or "").strip()):
            notfound.append(f"{title}|||{norm_issue(r[ii])}|||{str(r[bi] or '').strip()}")

    notfound = sorted(set(notfound))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(notfound, open(OUT, "w"), indent=0)
    print(f"Inventory: {os.path.basename(xlsx)}")
    print(f"Rows checked: {total:,}")
    print(f"NOT in GCD:   {len(notfound):,}  ({100*len(notfound)/total:.1f}%)")
    print(f"Wrote: {os.path.relpath(OUT, ROOT)}")
    print("Next: python3 brb.py --commit \"GCD not-found list\" --yes")


if __name__ == "__main__":
    main()
