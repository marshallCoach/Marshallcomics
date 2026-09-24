#!/usr/bin/env python3
"""brb_gcd_export_unknowns.py — turn the "Not in GCD" residual into a research
worklist.

Reads artifacts/comics-inventory/public/gcd_notfound.json (the books GCD can't
place, after cover-verified and cover-buy exclusions) and writes a CSV enriched
for a fast ComicVine / Fandom pass: publisher, year, on-sale date, ready-made
search links, and blank columns to fill (Fandom_Link, Volume, Notes).

Two buckets, GENUINE first:
  GENUINE — older books GCD lists under a different title (Giant-Size specials,
            FCBD editions, Milestone "Season One" reprints) → real research.
  RECENT  — published 2024+ → usually just GCD indexing lag; skim, don't sweat.

Read-only on the xlsx. Run on the Mac AFTER refreshing the list:
    python3 brb_gcd_notfound.py        # refresh the residual first
    python3 brb_gcd_export_unknowns.py # then export the worklist
Output: gcd_unknowns_review.csv  (open in Excel/Numbers/Sheets)
"""
import csv, glob, json, os, re, sys
from urllib.parse import quote_plus

ROOT = os.path.dirname(os.path.abspath(__file__))
NOTFOUND = os.path.join(ROOT, "artifacts/comics-inventory/public/gcd_notfound.json")
OUT = os.path.join(ROOT, "gcd_unknowns_review.csv")
RECENT_FROM = 2024  # year >= this is treated as GCD indexing lag, not a real gap


def _fn_key(f):
    m = re.search(r"_(\d{2})(\d{2})_(\d{2})(\d{2})", os.path.basename(f))
    if m:
        dd, mo, hh, mi = (int(x) for x in m.groups())
        return (1, mo, dd, hh, mi, os.path.getmtime(f))
    return (0, 0, 0, 0, 0, os.path.getmtime(f))


def latest_xlsx():
    m = [f for f in glob.glob(os.path.join(ROOT, "attached_assets/comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=_fn_key)


def norm_issue(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def year_of(pub_date, year):
    m = re.match(r"(\d{4})", str(pub_date or "")) or re.match(r"(\d{4})", str(year or ""))
    return int(m.group(1)) if m else None


def main():
    try:
        ids = json.load(open(NOTFOUND))
    except (OSError, ValueError):
        sys.exit(f"can't read {NOTFOUND} — run brb_gcd_notfound.py first")
    if len(ids) > 150:
        print(f"  ⚠ {len(ids)} entries — that looks stale. Run "
              f"python3 brb_gcd_notfound.py first, then re-run this.")

    import openpyxl
    xlsx = latest_xlsx()
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("✅ Clean Inventory"))
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0])

    def col(*names):
        for n in names:
            if n in H:
                return H.index(n)
        return None
    ci = {k: col(*v) for k, v in {
        "title": ("Title",), "issue": ("Issue #", "Issue"), "pub": ("Publisher",),
        "year": ("Year",), "box": ("Box #", "Box"), "pd": ("Publication Date", "Pub_Date"),
    }.items()}

    # index the inventory by (title.lower, issue-norm, box) for enrichment
    inv = {}
    for r in rows[1:]:
        t = str(r[ci["title"]] or "").strip()
        if not t:
            continue
        key = (t.lower(), norm_issue(r[ci["issue"]]), str(r[ci["box"]] or "").strip())
        inv[key] = {
            "Publisher": str(r[ci["pub"]] or "").strip() if ci["pub"] is not None else "",
            "Year": str(r[ci["year"]] or "").strip() if ci["year"] is not None else "",
            "Pub_Date": str(r[ci["pd"]] or "").strip() if ci["pd"] is not None else "",
        }

    out = []
    for cid in ids:
        parts = cid.split("|||")
        if len(parts) != 3:
            continue
        title, issue, box = parts
        meta = inv.get((title.lower(), norm_issue(issue), box.strip()), {})
        pub = meta.get("Publisher", "")
        yr = meta.get("Year", "")
        pd = meta.get("Pub_Date", "")
        y = year_of(pd, yr)
        bucket = "RECENT" if (y and y >= RECENT_FROM) else "GENUINE"
        iss = norm_issue(issue)
        q = f"{title} {pub} #{iss} {y or ''}".strip()
        fandom = f"https://www.google.com/search?q={quote_plus(q + ' site:fandom.com')}"
        cv = f"https://comicvine.gamespot.com/search/?q={quote_plus(title + ' ' + str(iss))}&indices%5B0%5D=issue"
        out.append({
            "bucket": bucket, "Title": title, "Issue": iss, "Box": box.strip(),
            "Publisher": pub, "Year": yr, "Pub_Date": pd,
            "fandom_search": fandom, "comicvine_search": cv,
            "Fandom_Link": "", "Volume": "", "Notes": "",
        })

    # GENUINE first, then by title/issue
    out.sort(key=lambda d: (d["bucket"] != "GENUINE", d["Title"].lower(),
                            float(re.match(r"[\d.]+", d["Issue"] or "0").group()) if re.match(r"[\d.]+", d["Issue"] or "") else 1e9))

    cols = ["bucket", "Title", "Issue", "Box", "Publisher", "Year", "Pub_Date",
            "fandom_search", "comicvine_search", "Fandom_Link", "Volume", "Notes"]
    with open(OUT, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out)

    g = sum(1 for d in out if d["bucket"] == "GENUINE")
    r = len(out) - g
    print(f"Inventory : {os.path.basename(xlsx)}")
    print(f"Exported  : {len(out)} books  →  {os.path.relpath(OUT, ROOT)}")
    print(f"  GENUINE (real research): {g}")
    print(f"  RECENT  (GCD lag, skim): {r}")
    print("\nWorkflow: research each GENUINE row (links are in the CSV), then drop")
    print("the Fandom link on the Data Fix page's 'Not in GCD' chip and export as")
    print("usual — brb_fix.py applies it and the cover clears the flag for good.")


if __name__ == "__main__":
    main()
