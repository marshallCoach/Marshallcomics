#!/usr/bin/env python3
"""brb_add_intake.py — append this week's new-comic pulls to the canonical xlsx.

Reads a simple CSV and appends one row per line to the newest
attached_assets/comics_inventory_*.xlsx, mapping by header name. Sets
Date Added = today. Leaves creators/values/cover blank for the GCD + Fandom
pipelines to fill later. Warns on same-box duplicates (Title+Issue+Year+Box).
Dry-run by default; --apply writes a NEW dated xlsx (never edits in place).

CSV columns (header row required, order-free, extras ignored):
    title,issue,year,volume,publisher,box,era
'volume' and 'era' may be blank. 'era' defaults to Modern.

    python3 brb_add_intake.py --csv intake_2026-09-15.csv
    python3 brb_add_intake.py --csv intake_2026-09-15.csv --apply
Then: python3 brb.py --commit "Weekly intake 2026-09-15" --yes
"""
import argparse, csv, datetime, glob, os, sys

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "attached_assets")


def newest_xlsx():
    m = [f for f in glob.glob(os.path.join(ASSETS, "comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=os.path.getmtime)


def ni(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", required=True, help="intake CSV (title,issue,year,volume,publisher,box,era)")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--restamp", action="store_true",
                    help="don't add rows; fix Date_Added marker on existing rows matching the CSV "
                         "(title+issue+box) so Recent Purchases picks them up")
    args = ap.parse_args()

    if not os.path.exists(args.csv):
        sys.exit(f"CSV not found: {args.csv}")
    rows = list(csv.DictReader(open(args.csv)))
    if not rows:
        sys.exit("CSV has no data rows")
    print(f"INTAKE: {os.path.basename(args.csv)}  ({len(rows)} rows)")

    import openpyxl
    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    def ci(n):
        return H.index(n) + 1 if n in H else None
    cT, cI, cY, cV, cP, cB, cE, cD = (ci("Title"), ci("Issue #"), ci("Year"), ci("Volume"),
                                      ci("Publisher"), ci("Box #"), ci("Era"), ci("Date Added"))
    cPD = ci("Publication Date")
    if cPD is None:
        cPD = ws.max_column + 1; ws.cell(1, cPD, "Publication Date")
    for name, col in (("Title", cT), ("Issue #", cI), ("Year", cY), ("Box #", cB)):
        if col is None:
            sys.exit(f"required column '{name}' not found in sheet")

    # existing keys for same-box dup warn (Title+Issue+Year+Box, per validator rule 2)
    existing = set()
    for r in range(2, ws.max_row + 1):
        t = str(ws.cell(r, cT).value or "").strip().lower()
        if t:
            existing.add((t, ni(ws.cell(r, cI).value), ni(ws.cell(r, cY).value), str(ws.cell(r, cB).value or "").strip()))

    today = datetime.date.today().isoformat()
    # Date_Added must carry the "new-comics intake" marker — RecentPurchases.tsx
    # filters on /new-comics intake/i, so a bare date never shows there.
    stamp = datetime.date.today().strftime("%B %d, %Y") + " (new-comics intake)"

    # --restamp: fix Date_Added on rows already added (match by title+issue+box), add nothing
    if args.restamp:
        idx = {}
        for r in range(2, ws.max_row + 1):
            t = str(ws.cell(r, cT).value or "").strip().lower()
            if t:
                idx.setdefault((t, ni(ws.cell(r, cI).value), str(ws.cell(r, cB).value or "").strip()), r)
        fixed = 0
        for rec in rows:
            k = ((rec.get("title") or "").strip().lower(), ni(rec.get("issue")), str(rec.get("box") or "").strip())
            r = idx.get(k)
            if r and cD:
                ws.cell(r, cD, stamp)
                ws.cell(r, cPD, today)  # today as publish date → shows on Release Timeline
                fixed += 1
                print(f"  restamp: {rec.get('title')} #{ni(rec.get('issue'))} Box {k[2]} -> Date_Added={stamp}, Pub_Date={today}")
            elif not r:
                print(f"  NOT FOUND (skip): {rec.get('title')} #{ni(rec.get('issue'))} Box {k[2]}")
        print(f"\n  Rows restamped: {fixed}")
        if not args.apply:
            print("  DRY RUN — nothing written. Re-run with --restamp --apply."); return
        out = os.path.join(ASSETS, f"comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
        wb.save(out)
        print(f"\n  WROTE: {os.path.basename(out)}  ({fixed} rows restamped)")
        print(f"  Next: python3 brb.py --commit \"Restamp intake Date_Added\" --yes")
        return

    added, dups = 0, []
    at = ws.max_row + 1
    for rec in rows:
        title = (rec.get("title") or "").strip()
        issue = ni(rec.get("issue"))
        year = str(rec.get("year") or "").strip()
        box = str(rec.get("box") or "").strip()
        if not title or not box:
            print(f"  SKIP (missing title/box): {rec}"); continue
        k = (title.lower(), issue, ni(year), box)
        if k in existing:
            dups.append(f"{title} #{issue} ({year}) Box {box}")
        ws.cell(at, cT, title)
        ws.cell(at, cI, issue)
        ws.cell(at, cY, year)
        if cP: ws.cell(at, cP, (rec.get("publisher") or "").strip())
        if cB: ws.cell(at, cB, box)
        if cV and (rec.get("volume") or "").strip():
            ws.cell(at, cV, (rec.get("volume") or "").strip())
        if cE: ws.cell(at, cE, (rec.get("era") or "Modern").strip())
        if cD: ws.cell(at, cD, stamp)
        if cPD: ws.cell(at, cPD, today)  # new pull = released ~today → Recent + Release Timeline
        print(f"  + {title} #{issue} ({year}) {rec.get('publisher','')} -> Box {box}"
              + (f"  Vol {rec.get('volume')}" if (rec.get('volume') or '').strip() else ""))
        existing.add(k); added += 1; at += 1

    if dups:
        print(f"\n  ⚠ {len(dups)} already in inventory (same Title+Issue+Year+Box) — added anyway, verify:")
        for d in dups:
            print(f"     {d}")

    print(f"\n  Rows to add: {added}")
    if not args.apply:
        print("  DRY RUN — nothing written. Re-run with --apply.")
        return
    out = os.path.join(ASSETS, f"comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
    wb.save(out)
    print(f"\n  WROTE: {os.path.basename(out)}  (+{added} rows)")
    print(f"  Next: python3 brb.py --commit \"Weekly intake {today}\" --yes")


if __name__ == "__main__":
    main()
