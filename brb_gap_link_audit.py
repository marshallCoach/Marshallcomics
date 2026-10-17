#!/usr/bin/env python3
"""brb_gap_link_audit.py — for every book still MISSING a cover (and, optionally,
still missing credits/date), check whether a Fandom link you already supplied
would fix it.

Answers: "of the books still on Data Fix / Missing Covers, how many already have a
link that should resolve them?" A link can fail to land for two reasons this
surfaces:
  1. CORRUPTED TITLE — the row's title is wrong (e.g. 'Low Slain' for 'Lois Lane'),
     so the cover (often already in covers.json under the real name) is unreachable.
     Fix: python3 brb_title_from_links.py --apply   (then the cover appears).
  2. LINK NEVER RE-APPLIED — the export was consumed; re-feeding it fills the gap.
     Fix: --emit writes a fresh data-fixes file to ~/Downloads for the next brb_all.

Read-only report by default. Cover presence mirrors the app's resolver exactly
(3-part Title|||Issue|||Vol, 2-part fallbacks, and the normalized single-url index).

    python3 brb_gap_link_audit.py                 # report
    python3 brb_gap_link_audit.py --credits        # also audit blank credits/date
    python3 brb_gap_link_audit.py --emit           # + write a re-apply data-fixes file
"""
import argparse, datetime, glob, json, os, re, sys, urllib.parse
from collections import defaultdict
import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "attached_assets")
DOWNLOADS = os.path.expanduser("~/Downloads")
PROCESSED = os.path.join(ROOT, "processed_exports")
COVERS = os.path.join(ROOT, "covers.json")


def _fn_key(f):
    m = re.search(r"_(\d{2})(\d{2})_(\d{2})(\d{2})", os.path.basename(f))
    if m:
        dd, mm, hh, mi = (int(x) for x in m.groups())
        return (1, mm, dd, hh, mi, os.path.getmtime(f))
    return (0, 0, 0, 0, 0, os.path.getmtime(f))


def newest_xlsx():
    cands = [f for f in glob.glob(os.path.join(ASSETS, "comics_inventory_*.xlsx"))
             if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not cands:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(cands, key=_fn_key)


# ── normalization mirroring CoverImage.tsx ───────────────────────────────────
def norm_title(t):
    t = str(t or "").lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def norm_issue(i):
    s = str(i or "").replace("#", "").strip()
    return re.sub(r"^0+(\d)", r"\1", s)


def ni(i):
    return re.sub(r"\.0$", "", str(i or "").replace("#", "").strip())


def comp_norm(t):  # for title-equality (ignore 'The'/punct)
    return re.sub(r"[^a-z0-9]+", "", re.sub(r"^the\s+", "", str(t or "").lower()))


def canonical_from_url(u):
    m = re.match(r"https?://[^/]+/wiki/(.+)", u or "")
    if not m:
        return None
    path = urllib.parse.unquote(m.group(1))
    vm = re.match(r"^(.*?)_Vol_\d+(?:_.+)?$", path)
    name = (vm.group(1) if vm else path)
    return name.replace("_", " ").strip() or None


def load_covers():
    cov = json.load(open(COVERS))
    have_exact = {k for k, v in cov.items() if isinstance(v, dict) and v.get("url")}
    groups = defaultdict(set)
    for k, v in cov.items():
        if not (isinstance(v, dict) and v.get("url")):
            continue
        p = k.split("|||")
        if len(p) < 2:
            continue
        groups[f"{norm_title(p[0])}|||{norm_issue(p[1])}"].add(v["url"])
    norm_idx = {nk for nk, urls in groups.items() if len(urls) == 1}
    # title+issue (any volume) presence — for "does a cover exist under the real name"
    by_title_issue = {f"{norm_title(p[0])}|{norm_issue(p[1])}"
                      for k, v in cov.items() if isinstance(v, dict) and v.get("url")
                      for p in [k.split("|||")] if len(p) >= 2}
    return have_exact, norm_idx, by_title_issue


def has_cover(title, issue, vol, have_exact, norm_idx):
    iss = str(issue)
    if f"{title}|||{iss}|||{vol}" in have_exact:   return True
    if f"{title}|||{iss}" in have_exact:           return True
    if f"{title}|||#{iss.lstrip('#')}" in have_exact: return True
    if f"{norm_title(title)}|||{norm_issue(iss)}" in norm_idx: return True
    return False


def link_files():
    seen, out = set(), []
    for d in (DOWNLOADS, PROCESSED):
        for f in glob.glob(os.path.join(d, "*.json")):
            if "data-fix" in os.path.basename(f).lower() and f not in seen:
                seen.add(f); out.append(f)
    return out


def load_links():
    """(key 'normcmp(title)|issue' -> {sheet,canon,url,issue,box}) from supplied links."""
    idx = {}
    for f in link_files():
        try:
            recs = json.load(open(f))
        except Exception:  # noqa: BLE001
            continue
        for rec in (recs if isinstance(recs, list) else []):
            if not isinstance(rec, dict):
                continue
            url = rec.get("value") or rec.get("url") or ""
            if "/wiki/" not in url:
                continue
            sheet = (rec.get("title") or "").strip()
            iss = ni(rec.get("issue"))
            if not sheet:
                continue
            idx[f"{comp_norm(sheet)}|{iss}"] = {
                "sheet": sheet, "canon": canonical_from_url(url) or sheet,
                "url": url, "issue": iss, "box": str(rec.get("box") or "").strip(),
            }
    return idx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--credits", action="store_true", help="also audit blank credits/date")
    ap.add_argument("--emit", action="store_true", help="write a re-apply data-fixes file")
    args = ap.parse_args()

    have_exact, norm_idx, cover_by_ti = load_covers()
    links = load_links()
    print(f"Supplied links: {len(links)} (from {len(link_files())} file(s))")

    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    def ci(n): return H.index(n) + 1 if n in H else None
    cT, cI, cV, cB = ci("Title"), ci("Issue #"), ci("Volume"), ci("Box #")
    cW, cA, cC, cPD = ci("Writer(s)"), ci("Artist(s)"), ci("Cover Artist"), ci("Publication Date")
    blank = lambda v: str(v or "").strip() in ("", "nan", "None")

    stats = defaultdict(int)
    rename_cover = []   # corrupted title AND a cover exists under the real name
    link_cover = []     # missing cover, link exists (title ok)
    emit = []
    for r in range(2, ws.max_row + 1):
        title = str(ws.cell(r, cT).value or "").strip()
        if not title:
            continue
        issue = str(ws.cell(r, cI).value or "").strip() if cI else ""
        vol = str(ws.cell(r, cV).value or "1").strip() if cV else "1"
        box = str(ws.cell(r, cB).value or "").strip() if cB else ""

        no_cover = not has_cover(title, issue, vol, have_exact, norm_idx)
        gaps = []
        if no_cover: gaps.append("cover")
        if args.credits:
            if cW and blank(ws.cell(r, cW).value): gaps.append("writer")
            if cA and blank(ws.cell(r, cA).value): gaps.append("artist")
            if cC and blank(ws.cell(r, cC).value): gaps.append("cover-artist")
            if cPD and blank(ws.cell(r, cPD).value): gaps.append("date")
        if not gaps:
            continue
        stats["gap_rows"] += 1
        if no_cover: stats["missing_cover"] += 1

        lk = links.get(f"{comp_norm(title)}|{ni(issue)}")
        if not lk:
            stats["gap_no_link"] += 1
            continue
        stats["gap_has_link"] += 1
        corrupted = comp_norm(lk["canon"]) != comp_norm(title)
        if corrupted:
            stats["link_corrupted_title"] += 1
            if no_cover and f"{norm_title(lk['canon'])}|{norm_issue(issue)}" in cover_by_ti:
                rename_cover.append((title, lk["canon"], issue))
        if no_cover and not corrupted:
            link_cover.append((title, issue, lk["url"]))
        # re-apply record uses the corrected (canonical) title so it matches post-rename
        emit.append({"title": lk["canon"], "issue": issue, "box": box, "value": lk["url"]})

    print("\n── GAP / LINK AUDIT ─────────────────────────────")
    print(f"  Gap rows (need a fix):        {stats['gap_rows']}")
    print(f"    · missing a cover:          {stats['missing_cover']}")
    print(f"  Gap rows WITH a supplied link:{stats['gap_has_link']}")
    print(f"    · title is corrupted:       {stats['link_corrupted_title']}  (fix: brb_title_from_links.py --apply)")
    print(f"  Gap rows with NO link:        {stats['gap_no_link']}  (need a new link)")

    if rename_cover:
        print(f"\n  ⤷ {len(rename_cover)} missing-cover book(s) whose cover ALREADY EXISTS under the real name")
        print(f"    (renaming the title instantly surfaces the cover):")
        for t, canon, iss in rename_cover[:40]:
            print(f"      {t!r} #{iss}  ->  {canon!r}")
    if link_cover:
        print(f"\n  ⤷ {len(link_cover)} missing-cover book(s) with a link (title ok) — re-apply to fetch:")
        for t, iss, _ in link_cover[:40]:
            print(f"      {t} #{iss}")

    if args.emit and emit:
        out = os.path.join(DOWNLOADS, f"data-fixes-reapply-{datetime.datetime.now():%Y%m%d_%H%M%S}.json")
        json.dump(emit, open(out, "w"), indent=2)
        print(f"\nWROTE {out}  ({len(emit)} records)")
        print("Order: 1) brb_title_from_links.py --apply   2) brb_all.py  (re-applies this file)")
    elif args.emit:
        print("\nNothing to emit.")


if __name__ == "__main__":
    main()
