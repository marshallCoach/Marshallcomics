#!/usr/bin/env python3
"""brb_title_from_links.py — REVIEW-ONLY title fixer from supplied Fandom links.

A data-fix link names a wiki page in its URL. Sometimes that page name reveals a
corrupted sheet title (voice-to-text 'Low Slain' for 'Lois Lane'). BUT most links
point to a legitimately differently-named page — a crossover, a one-shot, a
specific volume — so a URL name that differs is NOT proof the title is wrong.

An earlier version auto-applied every mismatch and wrongly renamed 500+ books.
This version NEVER auto-applies. It writes a review CSV of proposals (ranked, with
a similarity score and a flag), and you approve each one by hand. --apply then
renames ONLY the rows you marked, and carries their covers across.

    python3 brb_title_from_links.py            # write title_rename_review.csv
    # open the CSV, put Y in DECISION for real corruptions (or type a replacement
    # title to override); leave blank / N to skip.
    python3 brb_title_from_links.py --apply     # apply ONLY the approved rows
Then: python3 brb.py --commit "rename approved corrupted titles" --yes
"""
import argparse, csv, datetime, difflib, glob, json, os, re, sys, urllib.parse
import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "attached_assets")
DOWNLOADS = os.path.expanduser("~/Downloads")
PROCESSED = os.path.join(ROOT, "processed_exports")
COVERS = os.path.join(ROOT, "covers.json")
REVIEW = os.path.join(ROOT, "title_rename_review.csv")


def _fn_key(f):
    m = re.search(r"_(\d{2})(\d{2})_(\d{2})(\d{2})", os.path.basename(f))
    if m:
        dd, mm, hh, mi = (int(x) for x in m.groups())
        return (1, mm, dd, hh, mi, os.path.getmtime(f))
    return (0, 0, 0, 0, 0, os.path.getmtime(f))


def newest_xlsx():
    cands = [f for f in glob.glob(os.path.join(ASSETS, "comics_inventory_*.xlsx"))
             if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not cands:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(cands, key=_fn_key)


def ni(i):
    return re.sub(r"\.0$", "", str(i or "").replace("#", "").strip())


def normt(t):
    s = re.sub(r"^the\s+", "", str(t or "").lower())
    return re.sub(r"[^a-z0-9]+", "", s)


def canonical_from_url(u):
    m = re.match(r"https?://[^/]+/wiki/(.+)", u or "")
    if not m:
        return None
    path = urllib.parse.unquote(m.group(1)).split("?")[0].split("#")[0]  # drop query/anchor
    vm = re.match(r"^(.*?)_Vol_\d+(?:_.+)?$", path)
    name = (vm.group(1) if vm else path).replace("_", " ").strip()
    return name or None


# A canonical that still looks like a parse artifact or a different book, not a
# clean title — never auto-suggest these (they must be eyeballed).
_BOGUS = re.compile(r"(\?|#|file=|\bvol\b|\bvolume\b|\bissue\b|\(vol|\bannual\b|\bspecial\b|\d+\s*$)", re.I)


def link_files():
    seen, out = set(), []
    for d in (DOWNLOADS, PROCESSED):
        for f in glob.glob(os.path.join(d, "*.json")):
            if "data-fix" in os.path.basename(f).lower() and f not in seen:
                seen.add(f); out.append(f)
    return out


def collect_proposals():
    """One proposal per distinct sheet-title whose link names a different title."""
    prop = {}   # normt(sheet) -> dict
    for f in link_files():
        try:
            recs = json.load(open(f))
        except Exception:  # noqa: BLE001
            continue
        for rec in (recs if isinstance(recs, list) else []):
            if not isinstance(rec, dict):
                continue
            url = rec.get("value") or rec.get("url") or ""
            if "/wiki/" not in url:
                continue
            sheet = (rec.get("title") or "").strip()
            canon = canonical_from_url(url)
            if not sheet or not canon or normt(sheet) == normt(canon):
                continue
            sim = round(difflib.SequenceMatcher(None, normt(sheet), normt(canon)).ratio(), 2)
            prop[normt(sheet)] = {"old": sheet, "new": canon, "sim": sim,
                                  "box": str(rec.get("box") or "").strip(), "url": url}
    return prop


def write_review(prop, rows_by_title):
    out = []
    for k, p in prop.items():
        rows = rows_by_title.get(k, [])
        bogus = bool(_BOGUS.search(p["new"]))
        # "LIKELY" only for a close, clean, small-footprint match — still needs a Y.
        flag = ("LOOKS_BOGUS" if bogus else
                "LIKELY" if (p["sim"] >= 0.7 and len(rows) <= 5) else "NEEDS_REVIEW")
        out.append({"DECISION": "", "flag": flag, "similarity": p["sim"],
                    "rows_affected": len(rows), "old_title": p["old"],
                    "new_title": p["new"], "sample_url": p["url"]})
    order = {"LIKELY": 0, "NEEDS_REVIEW": 1, "LOOKS_BOGUS": 2}
    out.sort(key=lambda r: (order[r["flag"]], -r["similarity"], r["old_title"].lower()))
    with open(REVIEW, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["DECISION", "flag", "similarity", "rows_affected",
                                           "old_title", "new_title", "sample_url"])
        w.writeheader(); w.writerows(out)
    return out


def read_approvals():
    """{normt(old_title): replacement_title} for rows the user approved. DECISION:
    Y/yes -> use new_title; any other non-empty, non-N text -> that literal title."""
    if not os.path.exists(REVIEW):
        sys.exit(f"no {os.path.basename(REVIEW)} — run without --apply first, then approve rows.")
    appr = {}
    for r in csv.DictReader(open(REVIEW)):
        d = (r.get("DECISION") or "").strip()
        if not d or d.lower() in ("n", "no", "skip"):
            continue
        repl = r["new_title"] if d.lower() in ("y", "yes") else d
        appr[normt(r["old_title"])] = repl
    return appr


def remap_covers(pairs):
    if not os.path.exists(COVERS):
        return 0
    cov = json.load(open(COVERS))
    oldnorm = {normt(o): n for o, n in pairs}
    moved, out = 0, {}
    for k, v in cov.items():
        p = k.split("|||")
        n = oldnorm.get(normt(p[0])) if p else None
        if n:
            p[0] = n; nk = "|||".join(p)
            if nk not in cov:
                out[nk] = v; moved += 1; continue
        out.setdefault(k, v)
    if moved:
        json.dump(out, open(COVERS, "w"))
    return moved


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="apply ONLY rows marked in the review CSV")
    ap.add_argument("--no-text", action="store_true", help="don't scrub old title from Key/Notes")
    args = ap.parse_args()

    prop = collect_proposals()
    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    def ci(n): return H.index(n) + 1 if n in H else None
    cT, cKR, cSN = ci("Title"), ci("Key Issue — Why"), ci("Seller Notes / Variants / Caveats")

    rows_by_title = {}
    for r in range(2, ws.max_row + 1):
        rows_by_title.setdefault(normt(ws.cell(r, cT).value), []).append(r)

    if not args.apply:
        if not prop:
            print("No title mismatches in the supplied links. Nothing to review.")
            return
        rows = write_review(prop, rows_by_title)
        likely = sum(1 for r in rows if r["flag"] == "LIKELY")
        print(f"Wrote {os.path.basename(REVIEW)} — {len(rows)} proposal(s) "
              f"({likely} LIKELY, {len(rows)-likely} to scrutinize).")
        print("NOTHING was changed. Open the CSV, put Y in DECISION for real corruptions")
        print("(or type a replacement title to override), then re-run with --apply.")
        print("\nTop proposals:")
        for r in rows[:15]:
            print(f"  [{r['flag']:12s}] sim {r['similarity']:.2f}  x{r['rows_affected']:<3} "
                  f"{r['old_title']!r} -> {r['new_title']!r}")
        return

    appr = read_approvals()
    if not appr:
        print("No approved rows in the review CSV (DECISION column empty). Nothing to apply.")
        return
    applied_pairs, total = [], 0
    for r in range(2, ws.max_row + 1):
        k = normt(ws.cell(r, cT).value)
        if k in appr:
            new = appr[k]; old = str(ws.cell(r, cT).value)
            ws.cell(r, cT, new); total += 1
            applied_pairs.append((old, new))
            if not args.no_text:
                for c in (cKR, cSN):
                    if c and ws.cell(r, c).value:
                        ws.cell(r, c, re.sub(re.escape(old), new, str(ws.cell(r, c).value), flags=re.I))
    if not total:
        print("Approved titles matched no inventory rows. Nothing changed.")
        return
    out = os.path.join(ASSETS, f"comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
    wb.save(out)
    moved = remap_covers(list({(o, n) for o, n in applied_pairs}))
    print(f"Applied {len(set(k for k in appr))} approved rename(s) across {total} row(s).")
    print(f"WROTE {os.path.basename(out)}" + (f"  · moved {moved} cover key(s)." if moved else "."))
    print("Next: python3 brb.py --commit \"rename approved corrupted titles\" --yes")


if __name__ == "__main__":
    main()
