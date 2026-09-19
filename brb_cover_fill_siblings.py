#!/usr/bin/env python3
"""brb_cover_fill_siblings.py — fill a missing cover from a SIBLING copy.

Some books are missing a cover only because their cover lives under a different
volume key (a same title+issue copy elsewhere in the collection has it). Copying
that sibling cover is a free fill — BUT only when it's the same edition, or you
re-introduce wrong-era covers. So this is YEAR-GATED: it copies a sibling cover
to the missing book's volume-aware key only when the sibling cover's stored date
is within +/-1 year of the book's Year. Everything else is left alone.

Read-only on the xlsx (rows only); writes covers.json (+ public copy). Dry-run
by default.

    python3 brb_cover_fill_siblings.py
    python3 brb_cover_fill_siblings.py --apply
Then: python3 brb.py --commit "fill covers from year-matching siblings" --yes
"""
import argparse, glob, json, os, re, sys
from collections import defaultdict

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "attached_assets")
COVERS = os.path.join(ROOT, "covers.json")
PUBLIC = os.path.join(ROOT, "artifacts/comics-inventory/public/covers.json")
TOL = 1


def latest_xlsx():
    m = [f for f in glob.glob(os.path.join(ASSETS, "comics_inventory_*.xlsx"))
         if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not m:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(m, key=os.path.getmtime)


def ni(v):
    s = str(v or "").strip().lstrip("#")
    try:
        f = float(s); return str(int(f)) if f == int(f) else s
    except ValueError:
        return s


def nv(v):
    try:
        return str(int(float(str(v).strip())))
    except (ValueError, TypeError):
        return "1"


def yr(v):
    m = re.search(r"(19|20)\d{2}", str(v or ""))
    return int(m.group(0)) if m else None


def url_of(e):
    return (e.get("url") if isinstance(e, dict) else e) if e else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    cov = json.load(open(COVERS))
    # (title_lower, issue) -> list of (url, large, year) for keys that have a cover
    by_ti = defaultdict(list)
    for k, e in cov.items():
        u = url_of(e)
        if not u:
            continue
        p = k.split("|||")
        if len(p) < 2:
            continue
        large = e.get("large") if isinstance(e, dict) else u
        cy = yr(e.get("date")) if isinstance(e, dict) else None
        by_ti[(p[0].strip().lower(), ni(p[1]))].append((u, large or u, cy))

    import openpyxl
    src = latest_xlsx()
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("✅ Clean Inventory"))
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0])
    ti, ii, yi, vi = (H.index(c) for c in ("Title", "Issue #", "Year", "Volume"))

    filled, changes, seen = 0, [], set()
    for r in rows[1:]:
        t = str(r[ti] or "").strip()
        if not t:
            continue
        i = ni(r[ii]); vol = nv(r[vi]); ry = yr(r[yi])
        key = f"{t}|||{i}|||{vol}"
        dd = (t.lower(), i)
        # already resolvable by the app? (own 3-part or 2-part keys)
        if url_of(cov.get(key)) or url_of(cov.get(f"{t}|||{i}")) or url_of(cov.get(f"{t}|||#{i}")):
            continue
        if dd in seen:
            continue
        seen.add(dd)
        if not ry:
            continue
        # a sibling cover whose date matches this book's year
        best = next((s for s in by_ti.get(dd, []) if s[2] and abs(s[2] - ry) <= TOL), None)
        if not best:
            continue
        if args.apply:
            cov[key] = {"url": best[0], "large": best[1], "date": str(ry), "source": "sibling-fill"}
        filled += 1
        changes.append((t, i, vol, ry, best[2]))

    print(f"Inventory: {os.path.basename(src)}")
    print(f"Covers filled from year-matching siblings: {filled}")
    for t, i, v, ry, sy in changes[:40]:
        print(f"  {t} #{i} (Vol {v}, {ry}) <- sibling cover dated {sy}")
    if len(changes) > 40:
        print(f"  … and {len(changes) - 40} more")

    if not args.apply:
        print("\nDRY RUN — nothing written. Re-run with --apply.")
        return
    if not changes:
        print("\nNothing to fill.")
        return
    json.dump(cov, open(COVERS, "w"))
    if os.path.exists(PUBLIC):
        json.dump(cov, open(PUBLIC, "w"))
    print(f"\nWROTE covers.json (+ public) — {filled} covers filled.")
    print("Next: python3 brb.py --commit \"fill covers from year-matching siblings\" --yes")


if __name__ == "__main__":
    main()
