#!/usr/bin/env python3
"""brb_pubdate_cv.py — Comic Vine fallback for the Publication Date column.

Fills ONLY the rows GCD couldn't date (brb_pubdate_fill.py left blank). Queries
Comic Vine directly (free API — needs COMIC_VINE_API_KEY; NO Claude usage
credits) for each book's real on-sale date:
  1. store_date  (Comic Vine's actual on-sale date)  — preferred
  2. cover_date  (cover month)                        — fallback
Two-step match, same principle as the covers proxy: resolve the volume by
name + year + publisher, then read the issue inside it. Unmatched rows stay
blank — never guessed.

Runs on the Mac (that is where the key + internet live). Writes a NEW dated
xlsx; never edits the source. Fills blanks only unless --overwrite.

    python3 brb_pubdate_cv.py            # fill remaining blanks
    python3 brb_pubdate_cv.py --limit 50 # test on the first 50 blanks
Then: python3 brb.py --commit "Publication Date — Comic Vine fallback" --yes
"""
import argparse, datetime, glob, json, os, re, sys, time, urllib.parse, urllib.request
from collections import Counter

ROOT = os.path.dirname(os.path.abspath(__file__))
API_KEY = os.environ.get("COMIC_VINE_API_KEY", "").strip()
CV = "https://comicvine.gamespot.com/api"
UA = "BlackReadBrown-Comics/1.0"
COLNAME = "Publication Date"


def newest_xlsx():
    cands = [f for f in glob.glob(os.path.join(ROOT, "attached_assets/comics_inventory_*.xlsx"))
             if " copy" not in f and not os.path.basename(f).startswith("~$")]
    if not cands:
        sys.exit("no attached_assets/comics_inventory_*.xlsx found")
    return max(cands, key=os.path.getmtime)


def norm(t):
    return re.sub(r"[^a-z0-9]+", " ", str(t or "").lower()).strip()


def year_of(y):
    nums = [int(n) for n in re.findall(r"\d{4}", str(y or "")) if 1900 < int(n) < 2100]
    return nums[0] if nums else None


def cv_get(path, params, tries=5):
    """GET with polite backoff on Comic Vine's rate limiter."""
    params = {**params, "api_key": API_KEY, "format": "json"}
    url = f"{CV}/{path}/?{urllib.parse.urlencode(params)}"
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=30) as r:
                data = json.load(r)
            # status_code 107 = rate limited; 100 = invalid key
            if data.get("status_code") == 107:
                time.sleep(2 ** i * 5)
                continue
            if data.get("status_code") == 100:
                sys.exit("Comic Vine says the API key is invalid — check COMIC_VINE_API_KEY")
            return data
        except Exception as e:  # noqa: BLE001 — network hiccup: back off and retry
            if i == tries - 1:
                print(f"    ! giving up on request: {e}")
                return None
            time.sleep(2 ** i * 3)
    return None


def resolve_volume(cache, title, year, publisher):
    ck = (norm(title), year)
    if ck in cache:
        return cache[ck]
    data = cv_get("search", {"resources": "volume", "query": title,
                             "field_list": "id,name,start_year,publisher", "limit": "20"})
    best, best_score = None, -1e9
    for v in (data or {}).get("results", []) or []:
        score = 0
        if norm(v.get("name")) == norm(title):
            score += 10
        sy = v.get("start_year")
        try:
            sy = int(sy)
        except (TypeError, ValueError):
            sy = None
        if year and sy:
            score += 8 if abs(sy - year) <= 1 else (4 if abs(sy - year) <= 3 else -min(10, abs(sy - year)))
        pub = (v.get("publisher") or {}).get("name", "") if isinstance(v.get("publisher"), dict) else ""
        if publisher and pub:
            pl, pn = publisher.lower(), pub.lower()
            score += 5 if (pl in pn or pn in pl) else -3
        if score > best_score:
            best_score, best = score, v.get("id")
    cache[ck] = best
    time.sleep(0.4)
    return best


def issue_date(vol_id, issue):
    inum = str(issue).strip().lstrip("#")
    try:
        f = float(inum); inum = str(int(f)) if f == int(f) else inum
    except ValueError:
        pass
    data = cv_get("issues", {"filter": f"volume:{vol_id},issue_number:{inum}",
                             "field_list": "store_date,cover_date"})
    time.sleep(0.4)
    for r in (data or {}).get("results", []) or []:
        sd, cd = r.get("store_date"), r.get("cover_date")
        if sd and str(sd)[:4].isdigit():
            return sd, "store_date"
        if cd and str(cd)[:4].isdigit():
            return cd, "cover_date"
    return None, "no_date"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--overwrite", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="only process the first N blanks (testing)")
    args = ap.parse_args()
    if not API_KEY:
        sys.exit("COMIC_VINE_API_KEY not set. Add it to ~/.zshrc and: source ~/.zshrc")

    import openpyxl
    src = newest_xlsx()
    wb = openpyxl.load_workbook(src)
    ws = next(w for w in wb.worksheets if w.title.startswith("✅ Clean Inventory"))
    H = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    def ci(n):
        return H.index(n) + 1 if n in H else None
    cT, cI, cP, cY = ci("Title"), ci("Issue #"), ci("Publisher"), ci("Year")
    cPD = ci(COLNAME)
    if cPD is None:
        cPD = ws.max_column + 1
        ws.cell(1, cPD, COLNAME)
        print(f"added '{COLNAME}' column at index {cPD}")

    cache, stats, done = {}, Counter(), 0
    for r in range(2, ws.max_row + 1):
        title = str(ws.cell(r, cT).value or "").strip()
        if not title:
            continue
        existing = str(ws.cell(r, cPD).value or "").strip()
        if existing and not args.overwrite:
            continue
        stats["blanks"] += 1
        if args.limit and done >= args.limit:
            continue
        done += 1
        issue = str(ws.cell(r, cI).value or "").strip() if cI else ""
        pub = str(ws.cell(r, cP).value or "").strip() if cP else ""
        year = year_of(ws.cell(r, cY).value) if cY else None
        vol = resolve_volume(cache, title, year, pub)
        if not vol:
            stats["no_volume"] += 1
            continue
        date_str, kind = issue_date(vol, issue)
        stats[kind] += 1
        if date_str:
            ws.cell(r, cPD, date_str)
            stats["filled"] += 1
        if done % 25 == 0:
            print(f"  …{done} processed · {stats['filled']} filled")

    out = os.path.join(ROOT, f"attached_assets/comics_inventory_{datetime.datetime.now():%d%m_%H%M}.xlsx")
    wb.save(out)
    print(f"\nSOURCE: {os.path.basename(src)}")
    print(f"OUTPUT: {os.path.basename(out)}")
    print("\nCOVERAGE (blanks only)")
    for k in ("blanks", "filled", "store_date", "cover_date", "no_volume", "no_date"):
        if stats[k]:
            print(f"  {k:12s} {stats[k]}")
    print("\nNext: python3 brb.py --commit \"Publication Date — Comic Vine fallback\" --yes")


if __name__ == "__main__":
    main()
